import asyncio
import io
import logging
import re
from datetime import date, datetime
from typing import Any, Awaitable, Callable

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = ("pdf", "docx", "xlsx", "pptx", "txt", "md")

ImageUploader = Callable[[bytes, str, str], Awaitable[str]]


def clean_markdown(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\r\n", "\n")
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def escape_table_cell(value: str) -> str:
    return re.sub(r"\r?\n", " ", value.replace("|", "\\|")).strip()


def to_markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    max_cols = max((len(row) for row in rows), default=0)
    if max_cols == 0:
        return ""

    lines: list[str] = []
    for index, row in enumerate(rows):
        cells = [escape_table_cell(row[col] if col < len(row) else "") for col in range(max_cols)]
        lines.append("| " + " | ".join(cells) + " |")
        if index == 0:
            lines.append("| " + " | ".join(["---"] * max_cols) + " |")
    return "\n".join(lines) + "\n\n"


def get_extension(filename: str | None) -> str:
    if not filename or "." not in filename:
        return ""
    return filename.rsplit(".", 1)[-1].lower()


def decode_upload_filename(filename: str | None) -> str:
    if not filename:
        return ""
    try:
        decoded = filename.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return filename
    if "\ufffd" in decoded:
        return filename
    return decoded


def title_from_filename(filename: str | None) -> str:
    if not filename:
        return "未命名文档"
    index = filename.rfind(".")
    return filename[:index] if index > 0 else filename


def is_supported(extension: str) -> bool:
    return extension.lower() in SUPPORTED_EXTENSIONS


def supported_list() -> str:
    return ", ".join(SUPPORTED_EXTENSIONS)


# ---------------------------------------------------------------- DOCX


def _heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip().lower()
    if name in ("title", "标题"):
        return 1
    if name == "subtitle":
        return 2
    match = re.match(r"^(?:heading|标题)\s*([1-6])$", name)
    if match:
        return int(match.group(1))
    return None


def _list_prefix(style_name: str) -> str | None:
    name = (style_name or "").lower()
    if "list" not in name and "列表" not in name:
        return None
    if "number" in name or "编号" in name:
        return "1. "
    return "- "


def _paragraph_markdown(paragraph: Any) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style_name = paragraph.style.name if paragraph.style is not None else ""
    level = _heading_level(style_name)
    if level is not None:
        return f"{'#' * level} {text}"
    prefix = _list_prefix(style_name)
    if prefix is not None:
        return f"{prefix}{text}"
    return text


def parse_docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(_paragraph_markdown(Paragraph(child, document)))
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            parts.append(to_markdown_table(rows))
    return clean_markdown("\n\n".join(part for part in parts if part))


# ---------------------------------------------------------------- XLSX


def _xlsx_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"## {sheet.title}\n")

        rows: list[list[str]] = []
        max_cols = 0
        for row in sheet.iter_rows(values_only=True):
            if row is None:
                continue
            cells = [_xlsx_cell(value) for value in row]
            if not any(cell.strip() for cell in cells):
                continue
            max_cols = max(max_cols, len(cells))
            rows.append(cells)

        if max_cols == 0 or not rows:
            parts.append("\n")
            continue

        normalized = [row + [""] * (max_cols - len(row)) for row in rows]
        parts.append(to_markdown_table(normalized))
    workbook.close()
    return clean_markdown("\n".join(parts))


# ---------------------------------------------------------------- PPTX


def _pptx_paragraphs(text_frame: Any) -> list[str]:
    texts: list[str] = []
    for paragraph in text_frame.paragraphs:
        value = "".join(run.text for run in paragraph.runs).strip() or paragraph.text.strip()
        if value:
            texts.append(value)
    return texts


def parse_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"## 幻灯片 {index}\n")

        title_shape = slide.shapes.title
        title_text = ""
        if title_shape is not None and getattr(title_shape, "has_text_frame", False):
            title_text = title_shape.text.strip()

        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            rows = [
                [cell.text.strip() for cell in row.cells] for row in shape.table.rows
            ]
            parts.append(to_markdown_table(rows))

        if title_text:
            parts.append(f"### {title_text}\n")

        used = {title_text}
        body_lines: list[str] = []
        for shape in slide.shapes:
            if shape is title_shape or getattr(shape, "has_table", False):
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            for text in _pptx_paragraphs(shape.text_frame):
                if text in used:
                    continue
                body_lines.append(text)
        if body_lines:
            parts.append("\n".join(body_lines) + "\n\n")

    result = clean_markdown("\n".join(parts))
    if not result:
        raise ValueError("PPTX 提取结果为空")
    return result


# ---------------------------------------------------------------- PDF


def _sniff_image_content_type(data: bytes) -> str:
    if len(data) >= 3 and data[0] == 0xFF and data[1] == 0xD8 and data[2] == 0xFF:
        return "image/jpeg"
    if len(data) >= 8 and data[:4] == b"\x89PNG":
        return "image/png"
    if len(data) >= 4 and data[:4] == b"RIFF":
        return "image/webp"
    return "image/png"


async def parse_pdf(data: bytes, upload_image: ImageUploader | None = None) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        image_urls: list[str] = []

        if upload_image is not None:
            try:
                for image_index, image in enumerate(page.images):
                    payload = image.data
                    if not payload:
                        continue
                    content_type = _sniff_image_content_type(payload)
                    extension = "jpg" if content_type == "image/jpeg" else "png"
                    file_name = f"pdf_img_p{index}_{image_index}.{extension}"
                    try:
                        image_urls.append(
                            await upload_image(payload, file_name, content_type)
                        )
                    except Exception as exc:
                        logger.warning(
                            "PDF 图片上传失败: page=%s, name=%s, err=%s",
                            index,
                            getattr(image, "name", ""),
                            exc,
                        )
            except Exception as exc:
                logger.warning("PDF 图片提取失败，继续仅文本: page=%s, err=%s", index, exc)

        if text:
            parts.append(text)
        for url in image_urls:
            parts.append(f"![]({url})")
        if text or image_urls:
            parts.append("")

    return clean_markdown("\n\n".join(parts))


# ---------------------------------------------------------------- 入口


async def parse_file(
    data: bytes,
    original_name: str,
    upload_image: ImageUploader | None = None,
) -> str:
    extension = get_extension(original_name)
    if not is_supported(extension):
        raise ValueError(f"不支持的文件格式: {extension}，支持的格式: {supported_list()}")
    if not data:
        raise ValueError("文件内容为空，无法解析")

    if extension == "docx":
        result = await asyncio.to_thread(parse_docx, data)
    elif extension == "xlsx":
        result = await asyncio.to_thread(parse_xlsx, data)
    elif extension == "pptx":
        result = await asyncio.to_thread(parse_pptx, data)
    elif extension == "pdf":
        result = await parse_pdf(data, upload_image)
    else:
        result = data.decode("utf-8", errors="replace")

    if not result.strip():
        raise ValueError("文件解析结果为空，请确认文件包含可提取的文本内容")
    return result
